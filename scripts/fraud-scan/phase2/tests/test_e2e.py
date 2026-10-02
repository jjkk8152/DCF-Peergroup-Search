"""샘플 데이터 end-to-end (DART·LLM 호출 없음)."""
import io

import openpyxl

import evidence_trace
import run_sample
from export_excel import build_workbook


def test_end_to_end(tmp_path):
    out = run_sample.run(str(tmp_path / "e2e.db"), tmp_path / "e2e.xlsx", quiet=True)
    conn, eid, uid, s = out["conn"], out["engagement_id"], out["upload_id"], out["summary"]

    assert s["benchmark_companies_selected"] == 4
    assert s["relevant_fraud_cases"] == 6 and s["deficiency_disclosures"] == 1 and s["keyword_hits_excluded"] == 1
    assert s["rcm_controls_reviewed"] == 8
    assert sum(s["coverage"].values()) == s["fraud_scenarios_identified"] == 8
    assert set(s["coverage"]) <= {"Covered", "Partially Covered", "Not Covered", "Cannot Determine"}

    areas = evidence_trace.attention_areas(conn, eid, uid)
    assert areas[0]["scenario_id"] == "FS004" and areas[0]["coverage"] == "Not Covered"

    # 드릴다운: 시나리오 → 사례 → 공시 → 원문 근거 → 분류 → RCM 통제 → coverage → 감사절차
    t = evidence_trace.trace(conn, eid, uid, "FS003")
    assert t["cases"] and all(c["evidence"] and c["filing"]["source_url"] for c in t["cases"])
    assert all(c["classification"]["rationale"] for c in t["cases"])
    assert t["controls"] and t["controls"][0]["why_matched"]
    assert t["coverage"]["coverage"] == "Partially Covered"
    assert any("Contradictory Information" in i for i in t["coverage"]["information_needed"])
    a = t["audit_consideration"]
    assert a["conclusion"] == "Auditor judgment required"
    assert "deficiency" in a["control_coverage_gap"] and "자동으로 확정하지 않음" in a["audit_response_consideration"]
    assert "OTP 또는 공동인증서의 실제 보관자는 누구인가?" in " ".join(a["followup_questions"])

    wb = openpyxl.load_workbook(io.BytesIO(build_workbook(conn, eid, uid)))
    expected = ["1 Engagement Profile", "2 Selected Peers", "3 Fraud Cases", "4 Fraud Scenarios", "5 RCM Raw", "6 Scenario-Control Mapping",
                "7 Coverage Assessment", "8 Audit Considerations", "9 Source Evidence", "10 Review Log"]
    assert wb.sheetnames[:10] == expected
    for name in ["3 Fraud Cases", "6 Scenario-Control Mapping", "7 Coverage Assessment", "8 Audit Considerations"]:
        headers = [c.value for c in wb[name][1]]
        for h in ["AI Generated", "Reviewer", "Review Status", "Reviewer Comment"]:
            assert h in headers, (name, h)
    assert wb["7 Coverage Assessment"].max_row == 9
    assert wb["9 Source Evidence"].max_row > 7
    rcm_headers = [c.value for c in wb["5 RCM Raw"][1]]
    assert "통제 설명" in rcm_headers  # 원본 컬럼명 그대로
