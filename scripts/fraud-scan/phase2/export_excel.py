"""감사조서 초안용 Excel export (요구사항 20).

시트: 1 Engagement Profile · 2 Selected Peers · 3 Fraud Cases · 4 Fraud Scenarios · 5 RCM Raw ·
      6 Scenario-Control Mapping · 7 Coverage Assessment · 8 Audit Considerations · 9 Source Evidence · 10 Review Log
      (+ 11 Contradictions: 상충 정보 모듈 결과)
자동 생성 결과 시트에는 AI Generated / Reviewer / Review Status / Reviewer Comment 컬럼을 둔다.
Review Status 는 드롭다운(Not Reviewed/Accepted/Modified/Rejected) — 최종 판단은 reviewer 가 수행.
"""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import audit_response_engine
import coverage_engine
import fraud_case_engine
import fraud_scenario_engine
from db import REVIEW_STATUSES, rows, unj
from engagement import load_engagement
from guardrails import DISCLAIMER
from rcm_parser import load_rcm
from scenario_lib import scenario as get_scenario

REVIEW_HEADERS = ["AI Generated", "Reviewer", "Review Status", "Reviewer Comment"]


def _review(r: dict) -> dict:
    return {
        "AI Generated": "Yes" if r.get("ai_generated", 1) else "No",
        "Reviewer": r.get("reviewer") or "",
        "Review Status": r.get("review_status") or "Not Reviewed",
        "Reviewer Comment": r.get("reviewer_comment") or "",
    }


def _list(v) -> str:
    if isinstance(v, list):
        return "\n".join(f"- {x}" if not isinstance(x, dict) else f"- {x}" for x in v)
    return "" if v is None else str(v)


def build_workbook(conn, engagement_id: int, upload_id: int | None) -> bytes:
    p = load_engagement(conn, engagement_id)
    sheets: dict[str, pd.DataFrame] = {}

    sheets["1 Engagement Profile"] = pd.DataFrame([
        ("Company", p.company_name), ("Industry", p.industry), ("Subindustry", p.subindustry), ("KSIC", ", ".join(p.ksic_codes)),
        ("Market", p.market), ("Size basis", p.size_basis), ("Size (KRW)", p.size_value), ("Overseas subsidiary", p.has_overseas_subsidiary),
        ("Period", f"{p.period_from} ~ {p.period_to}"), ("Business description", p.business_description),
        ("Treasury features", p.treasury_features), ("Disclaimer", DISCLAIMER),
    ], columns=["Item", "Value"])

    peers = rows(conn, "SELECT * FROM peer_candidates WHERE engagement_id=? ORDER BY selected DESC, similarity_score DESC", [engagement_id])
    sheets["2 Selected Peers"] = pd.DataFrame([{
        "Selected (Auditor)": "Yes" if r["selected"] else "No", "Selected By": r["selected_by"], "Company": r["company_name"],
        "corp_code": r["corp_code"], "Stock code": r["stock_code"], "KSIC": r["industry_code"], "Industry": r["industry_name"],
        "Market": r["market"], "Similarity score": r["similarity_score"],
        "Selection reasons": _list((unj(r["reasons"], {}) or {}).get("similar")), "Differences": _list((unj(r["reasons"], {}) or {}).get("different")),
        "Score breakdown": r["score_breakdown"], **_review(r),
    } for r in peers])

    cases = fraud_case_engine.load_cases(conn, engagement_id, kinds=("fraud_incident", "control_deficiency_disclosure", "keyword_hit"))
    links = rows(conn, "SELECT * FROM case_scenarios WHERE engagement_id=?", [engagement_id])
    sid_by_case: dict[str, list[str]] = {}
    for l in links:
        sid_by_case.setdefault(l["case_id"], []).append(l["scenario_id"])
    sheets["3 Fraud Cases"] = pd.DataFrame([{
        "case_id": c["case_id"], "Case kind": c["case_kind"], "Basis": c["case_basis"], "Company": c["company_name"], "corp_code": c["corp_code"],
        "Industry": c["industry"], "rcept_no": c["rcept_no"], "Filing date": c["filing_date"], "Report": c["report_name"],
        "Fraud type": c["fraud_type"], "Fraud scheme": c["fraud_scheme"], "Actor": c["actor"], "Actor position": c["actor_position"],
        "Period of fraud": c["period_of_fraud"], "Amount (KRW)": c["amount"], "Affected accounts": _list(c["affected_accounts"]),
        "Cash mechanism": c["cash_mechanism"], "Transaction type": c["transaction_type"],
        "Related party involved": {None: "Not disclosed", 1: "Yes", 0: "No"}.get(c["related_party_involved"], "Not disclosed"),
        "Control weakness": _list(c["control_weakness"]), "Circumvention": _list(c["control_circumvention_method"]),
        "Detection method": c["detection_method"], "Remediation": _list(c["remediation"]), "Audit relevance": c["audit_relevance"],
        "Scenarios": ", ".join(sid_by_case.get(c["case_id"], [])), "Confidence": c["confidence"], "Extraction": c["extraction_method"],
        "Source URL": c["source_url"], **_review(c),
    } for c in cases])

    clusters = fraud_scenario_engine.cluster(conn, engagement_id)
    scen_rows = [{
        "Scenario ID": c["scenario_id"], "Title": f"{c['title']} ({c['title_ko']})", "Status": "Fixed Taxonomy", "Description": c["description"],
        "Mechanism": c["mechanism"], "Accounts": ", ".join(c["accounts"]), "Processes": ", ".join(c["processes"]), "Assertions": ", ".join(c["assertions"]),
        "Peer incident cases": c["peer_incident_count"], "Deficiency disclosures": c["peer_deficiency_disclosure_count"],
        "Companies": ", ".join(c["companies"]), "Filing period": c["filing_period"], "Case IDs": ", ".join(c["case_ids"]),
        "Note": c["frequency_note"], "AI Generated": "Yes", "Reviewer": "", "Review Status": "Not Reviewed", "Reviewer Comment": "",
    } for c in clusters if c["case_ids"]]
    for n in fraud_scenario_engine.new_scenario_candidates(conn, engagement_id):
        scen_rows.append({"Scenario ID": f"NEW-{n['id']}", "Title": n["proposed_title"], "Status": n["status"], "Mechanism": n["mechanism"],
                          "Case IDs": ", ".join(n["case_ids"]), "Note": "신규 시나리오 후보 — 감사인 검토 전 확정 아님", **_review(n)})
    sheets["4 Fraud Scenarios"] = pd.DataFrame(scen_rows)

    rcm = load_rcm(conn, upload_id) if upload_id else []
    sheets["5 RCM Raw"] = pd.DataFrame([{"row_index": r["row_index"], **r["raw"]} for r in rcm])

    maps, covs, auds = [], [], []
    if upload_id:
        for m in rows(conn, """SELECT m.*, r.control_id, r.control_description, r.risk_id, r.risk_description FROM scenario_control_matches m
                               JOIN rcm_rows r ON r.upload_id=m.upload_id AND r.row_index=m.row_index
                               WHERE m.engagement_id=? AND m.upload_id=? ORDER BY m.scenario_id, m.rank""", [engagement_id, upload_id]):
            maps.append({"Scenario ID": m["scenario_id"], "Rank": m["rank"], "RCM row": m["row_index"], "Risk ID": m["risk_id"],
                         "Risk description": m["risk_description"], "Control ID": m["control_id"], "Control description": m["control_description"],
                         "Match score": m["match_score"], "Matched elements": _list([f"{e['element_id']} {e['label']}: \"{e['quote']}\"" for e in unj(m["matched_elements"], [])]),
                         "Why matched": m["why_matched"], "Vague description": "Yes" if m["is_vague"] else "No",
                         "Vague reasons": _list(unj(m["vague_reasons"], [])), "Factor scores": m["factor_scores"], **_review(m)})
        for r in coverage_engine.load_results(conn, engagement_id, upload_id):
            covs.append({"Scenario ID": r["scenario_id"], "Scenario": r["scenario_title"], "Coverage (AI suggestion)": r["coverage"],
                         "Rationale": r["rationale"], "Covered elements": _list(r["covered_elements"]), "Uncovered elements": _list(r["uncovered_elements"]),
                         "Possible circumvention": _list(r["possible_circumvention"]), "Information needed": _list(r["information_needed"]),
                         "Matched controls": ", ".join(m["control_id"] for m in r["matched_controls"]), "External cases": ", ".join(r["external_cases"]),
                         "Design vs operating": r["design_vs_operating_note"], "Confidence": r["confidence"], "Method": r["method"],
                         "LLM notes": r["llm_notes"], **_review(r)})
        for a in audit_response_engine.load(conn, engagement_id, upload_id):
            auds.append({"Scenario ID": a["scenario_id"], "Coverage": a["coverage"], "Why this matters": a["why_this_matters"],
                         "Control Coverage Gap": a["control_coverage_gap"], "Audit Response Consideration": a["audit_response_consideration"],
                         "Auditor Follow-up Questions": _list(a["followup_questions"]), "Suggested Audit Procedures": _list(a["suggested_procedures"]),
                         "Potential Evidence": _list(a["potential_evidence"]), "Possible Effect on Risk Assessment": a["risk_assessment_effect"],
                         "Conclusion": a["conclusion"], **_review(a)})
    sheets["6 Scenario-Control Mapping"] = pd.DataFrame(maps)
    sheets["7 Coverage Assessment"] = pd.DataFrame(covs)
    sheets["8 Audit Considerations"] = pd.DataFrame(auds)

    ev = []
    for c in cases:
        for i, e in enumerate(c["source_evidence"] or []):
            ev.append({"case_id": c["case_id"], "Evidence #": i + 1, "Company": c["company_name"], "rcept_no": c["rcept_no"],
                       "Filing date": c["filing_date"], "Report": c["report_name"], "Section": e.get("section"),
                       "Location": e.get("page_or_location"), "Evidence text (verbatim)": e.get("text"),
                       "Scenarios": ", ".join(sid_by_case.get(c["case_id"], [])), "Source URL": c["source_url"]})
    sheets["9 Source Evidence"] = pd.DataFrame(ev)
    sheets["10 Review Log"] = pd.DataFrame(rows(conn, "SELECT at, entity, entity_key, action, reviewer, review_status, comment FROM review_log WHERE engagement_id=? OR engagement_id IS NULL ORDER BY id", [engagement_id]))
    sheets["11 Contradictions"] = pd.DataFrame([{
        "Topic": c["topic"], "Source 1": c["source_1"], "Statement 1": c["statement_1"], "Source 2": c["source_2"], "Statement 2": c["statement_2"],
        "Status": c["status"], "Auditor follow-up": c["auditor_follow_up"], **_review(c),
    } for c in rows(conn, "SELECT * FROM contradictions WHERE engagement_id=?", [engagement_id])])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            (df if not df.empty else pd.DataFrame({"(no data)": []})).to_excel(xw, sheet_name=name[:31], index=False)
        wb = xw.book
        header_fill = PatternFill("solid", fgColor="DDE3EA")
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for col in ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)):
                width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = max(10, min(60, width + 2))
                for c in col[1:]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            headers = [c.value for c in ws[1]]
            if "Review Status" in headers and ws.max_row >= 2:
                col = get_column_letter(headers.index("Review Status") + 1)
                dv = DataValidation(type="list", formula1='"' + ",".join(REVIEW_STATUSES) + '"', allow_blank=False)
                ws.add_data_validation(dv)
                dv.add(f"{col}2:{col}{ws.max_row}")
    return buf.getvalue()


_ = get_scenario  # (시트 4는 cluster 결과 사용)
